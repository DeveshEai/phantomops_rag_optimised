"""
tools/excel_creator.py — Excel Report Generator

Creates a styled multi-sheet .xlsx report from structured analysis (NOT a
format conversion). Used when the user asks for analysis delivered as Excel.

The agent supplies sheets as JSON: each sheet has a name, headers, and rows.
"""

import os
from datetime import datetime
from langchain_core.tools import tool
from core.config import PROPOSALS_OUTPUT_DIR

_NAVY  = "1F3355"
_WHITE = "FFFFFF"
_LIGHT = "E8EEF8"


@tool
def create_excel_report(
    title: str,
    sheets: list,
    model_label: str = "",
    task_number: int = 0,
) -> str:
    """
    Create a styled Excel (.xlsx) report from structured data.
    Use this when the user asks for analysis or a report delivered as an Excel/spreadsheet.

    Parameters:
      title        : Report title (used for the filename).
      sheets       : List of sheet objects, each:
                       {"name": "Summary",
                        "headers": ["Metric", "Value"],
                        "rows": [["Total Employees", "45"], ["Avg Score", "7.8"]]}
                     Provide multiple sheets — e.g. a Summary sheet plus
                     Top Performers, Rankings, Department Breakdown, etc.
      model_label  : Model that generated this (for the filename).
      task_number  : Task number (for the filename).

    Returns the absolute path to the saved .xlsx file, or an error message.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        if not sheets:
            return "Error: no sheets provided for the Excel report."

        wb = Workbook()
        wb.remove(wb.active)

        header_fill = PatternFill("solid", fgColor=_NAVY)
        header_font = Font(bold=True, color=_WHITE, size=11, name="Calibri")
        alt_fill    = PatternFill("solid", fgColor=_LIGHT)
        body_font   = Font(size=10, name="Calibri")
        thin        = Side(style="thin", color="D0D7E5")
        border      = Border(left=thin, right=thin, top=thin, bottom=thin)

        used_names: set = set()

        def _safe_sheet_name(raw: str, idx: int) -> str:
            # Excel forbids : \ / ? * [ ] and titles must be 1-31 chars, unique.
            import re as _re
            name = _re.sub(r"[:\\/?*\[\]]", " ", str(raw or "")).strip()
            name = name[:31] or f"Sheet{idx}"
            base, n = name, 1
            while name.lower() in used_names:
                suffix = f" ({n})"
                name = base[:31 - len(suffix)] + suffix
                n += 1
            used_names.add(name.lower())
            return name

        for s_idx, sheet in enumerate(sheets, 1):
            name    = _safe_sheet_name(sheet.get("name", f"Sheet{s_idx}"), s_idx)
            headers = sheet.get("headers", [])
            rows    = sheet.get("rows", [])
            ws      = wb.create_sheet(title=name)

            # Header row
            if headers:
                ws.append([str(h) for h in headers])
                for col_idx in range(1, len(headers) + 1):
                    c = ws.cell(row=1, column=col_idx)
                    c.fill      = header_fill
                    c.font      = header_font
                    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    c.border    = border

            # Data rows with alternating shading
            for r_idx, row in enumerate(rows, start=2):
                ws.append([("" if v is None else str(v)) for v in row])
                for col_idx in range(1, len(row) + 1):
                    c = ws.cell(row=r_idx, column=col_idx)
                    c.font   = body_font
                    c.border = border
                    if r_idx % 2 == 0:
                        c.fill = alt_fill

            # Auto-fit column widths (capped)
            ncols = max([len(headers)] + [len(r) for r in rows] + [1])
            for col_idx in range(1, ncols + 1):
                max_len = 0
                for cell in ws[get_column_letter(col_idx)]:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 3, 12), 50)

            ws.freeze_panes = "A2"  # keep header visible on scroll

        os.makedirs(PROPOSALS_OUTPUT_DIR, exist_ok=True)
        ts           = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_title   = "".join(c if c.isalnum() or c in " _-" else "_" for c in title)[:50].strip()
        model_prefix = f"{model_label.upper()}_" if model_label else ""
        task_tag     = f"T{task_number}_" if task_number else ""
        out = os.path.join(PROPOSALS_OUTPUT_DIR, f"{model_prefix}{task_tag}{safe_title}_{ts}.xlsx")
        wb.save(out)
        print(f"[XLSX] Saved Excel report ({len(sheets)} sheet(s)) -> {out}")
        return out

    except Exception as e:
        return f"Error creating Excel report: {e}"


@tool
def create_csv_report(
    title: str,
    headers: list,
    rows: list,
    model_label: str = "",
    task_number: int = 0,
) -> str:
    """
    Create a CSV (.csv) report from a single flat table.
    Use when the user asks for the result/analysis delivered as a CSV file.
    CSV holds ONE table only — pick the most important data for the columns/rows.

    Parameters:
      headers      : list of column names.
      rows         : list of rows, each a list of cell values.
      model_label  : model that generated this (for the filename).
      task_number  : task number (for the filename).

    Returns the absolute path to the saved .csv file, or an error message.
    """
    try:
        import csv

        os.makedirs(PROPOSALS_OUTPUT_DIR, exist_ok=True)
        ts           = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_title   = "".join(c if c.isalnum() or c in " _-" else "_" for c in title)[:50].strip()
        model_prefix = f"{model_label.upper()}_" if model_label else ""
        task_tag     = f"T{task_number}_" if task_number else ""
        out = os.path.join(PROPOSALS_OUTPUT_DIR, f"{model_prefix}{task_tag}{safe_title}_{ts}.csv")

        with open(out, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            if headers:
                writer.writerow([str(h) for h in headers])
            for row in rows or []:
                writer.writerow([("" if v is None else str(v)) for v in row])

        print(f"[CSV] Saved CSV report ({len(rows or [])} row(s)) -> {out}")
        return out

    except Exception as e:
        return f"Error creating CSV report: {e}"


def _parse_markdown_tables(content: str) -> list:
    """
    Parse markdown into a list of tables: [{"name", "headers", "rows"}].
    A '## Heading' line just before a table becomes that table's sheet name.
    """
    import re as _re
    lines   = (content or "").replace("\r\n", "\n").split("\n")
    tables  = []
    cur_name = None
    pending_heading = "Sheet1"
    buf = []

    def _flush():
        nonlocal buf, pending_heading
        if not buf:
            return
        rows = []
        for ln in buf:
            cells = [c.strip() for c in ln.strip().strip("|").split("|")]
            if all((c and set(c) <= set("-: ")) for c in cells):
                continue  # separator row
            rows.append(cells)
        buf = []
        if rows:
            tables.append({
                "name":    pending_heading or f"Sheet{len(tables)+1}",
                "headers": rows[0],
                "rows":    rows[1:],
            })

    heading_re = _re.compile(r"^#{1,4}\s+(.*)$")
    for raw in lines:
        line = raw.strip()
        if line.startswith("|"):
            buf.append(line)
        else:
            _flush()
            h = heading_re.match(line)
            if h:
                pending_heading = h.group(1).strip()[:31]
    _flush()
    return tables


@tool
def create_spreadsheet(title: str, content: str, format: str = "xlsx") -> str:
    """
    Create an Excel or CSV file from markdown tables.

    Write one or more markdown tables in `content`. Put a '## Sheet Name' line before
    each table to name its sheet (Excel). Example:
      ## Summary
      | Metric | Value |
      |--------|-------|
      | Total  | 45    |
      ## Top Performers
      | Rank | Name | Score |
      |------|------|-------|
      | 1    | Rahul| 9.8   |

    Use this when the user wants the result as a spreadsheet/Excel/CSV.

    Args:
        title:   File title (used for the filename).
        content: One or more markdown tables (with optional '## name' headings).
        format:  "xlsx" (default, supports multiple sheets) or "csv" (first table only).

    Returns the absolute path to the created file, or an error message.
    """
    try:
        tables = _parse_markdown_tables(content)
        if not tables:
            return "Error: no markdown tables found in content for the spreadsheet."

        fmt = (format or "xlsx").lower().strip()

        if fmt == "csv":
            t = tables[0]
            return create_csv_report.func(
                title=title, headers=t["headers"], rows=t["rows"],
            ) if hasattr(create_csv_report, "func") else create_csv_report.invoke({
                "title": title, "headers": t["headers"], "rows": t["rows"],
            })

        # xlsx — reuse the styled multi-sheet generator
        sheets = [{"name": t["name"], "headers": t["headers"], "rows": t["rows"]} for t in tables]
        return create_excel_report.invoke({"title": title, "sheets": sheets})

    except Exception as e:
        return f"Error creating spreadsheet: {e}"
